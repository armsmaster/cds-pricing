from __future__ import annotations

import json
import os
from collections.abc import Awaitable, Callable, Iterator
from datetime import date
from pathlib import Path
from typing import Any

from fastapi import Depends, FastAPI, HTTPException, Query, Request, Response
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from sqlalchemy.orm import Session

from app.backend import cds_service, credit_service, data_service, excel, repositories, service
from app.backend import repositories as repo
from app.backend.db import get_session
from app.backend.moex import MoexClient
from app.backend.schemas import (
    AddBondRequest,
    BootstrapRequest,
    CreditCurveRequest,
    IssuerIn,
    PriceOverride,
    QuotesRequest,
    RecoveryPatch,
)

_FRONTEND = Path(__file__).resolve().parent.parent / "frontend"

_BUNDLE = "/etc/ssl/certs/ca-certificates.crt"

app = FastAPI(title="CRD — Credit Derivatives Pricing")


def get_moex() -> Iterator[MoexClient]:
    raw = os.environ.get("PCDS_MOEX_VERIFY")
    if raw is not None:
        verify_disable = raw.lower() in ("0", "false", "no")
        verify: bool | str = not verify_disable
    else:
        verify = _BUNDLE if os.path.isfile(_BUNDLE) else True
    client = MoexClient(verify=verify)
    try:
        yield client
    finally:
        client.close()


@app.middleware("http")
async def no_cache_static(
    request: Request, call_next: Callable[[Request], Awaitable[Response]]
) -> Response:
    response = await call_next(request)
    if not request.url.path.startswith("/api"):
        response.headers["Cache-Control"] = "no-cache, must-revalidate"
    return response


# --- rate curve (OIS) -----------------------------------------------------


@app.post("/api/quotes")
def quotes(
    request: QuotesRequest, session: Session = Depends(get_session)
) -> dict[str, Any]:
    try:
        gen = service._get_generator(session)
        return service.preview(request.quotes, request.trade_date, gen)
    except (KeyError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/api/bootstrap")
def bootstrap_endpoint(
    request: BootstrapRequest, session: Session = Depends(get_session)
) -> dict[str, Any]:
    try:
        gen = service._get_generator(session)
        return service.run_bootstrap(
            request.quotes,
            request.trade_date,
            request.max_adjustment_bps,
            request.smoothing_pct,
            gen,
        )
    except (KeyError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/api/rate-curves")
def save_rate_curve(
    request: BootstrapRequest, session: Session = Depends(get_session)
) -> dict[str, Any]:
    try:
        result = service.compute(
            request.quotes, request.trade_date, request.max_adjustment_bps, request.smoothing_pct
        )
    except (KeyError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    record = repositories.save_rate_curve(
        session,
        result.rate_index,
        result.curve,
        trade_date=result.trade_date,
        zcyc_json=None,
        max_adjustment_bps=request.max_adjustment_bps,
        average_adjustment_bps=result.average_adjustment_bps,
    )
    return credit_service.rate_curve_dict(record)


@app.get("/api/rate-curves")
def list_rate_curves(session: Session = Depends(get_session)) -> list[dict[str, Any]]:
    return [credit_service.rate_curve_dict(r) for r in repositories.list_rate_curves(session)]


@app.delete("/api/rate-curves/{curve_id}")
def delete_rate_curve(
    curve_id: int, session: Session = Depends(get_session)
) -> dict[str, str]:
    repositories.delete_rate_curve(session, curve_id)
    return {"status": "deleted"}


# --- issuers --------------------------------------------------------------


@app.post("/api/issuers")
def create_issuer(
    request: IssuerIn, session: Session = Depends(get_session)
) -> dict[str, Any]:
    issuer = repositories.create_issuer(
        session, request.name, request.recovery_rate, request.moex_emitent_id
    )
    return credit_service.issuer_dict(issuer)


@app.get("/api/issuers")
def list_issuers(session: Session = Depends(get_session)) -> list[dict[str, Any]]:
    return [credit_service.issuer_dict(i) for i in repositories.list_issuers(session)]


@app.patch("/api/issuers/{issuer_id}")
def patch_issuer(
    issuer_id: int, request: RecoveryPatch, session: Session = Depends(get_session)
) -> dict[str, Any]:
    try:
        issuer = repositories.set_recovery_rate(session, issuer_id, request.recovery_rate)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    return credit_service.issuer_dict(issuer)


@app.delete("/api/issuers/{issuer_id}")
def delete_issuer(
    issuer_id: int, session: Session = Depends(get_session)
) -> dict[str, str]:
    repositories.delete_issuer(session, issuer_id)
    return {"status": "deleted"}


# --- bonds ----------------------------------------------------------------


@app.get("/api/moex/search")
def moex_search(
    q: str = Query(min_length=1), client: MoexClient = Depends(get_moex)
) -> list[dict[str, Any]]:
    try:
        return client.search(q)
    except Exception as exc:  # noqa: BLE001 - surface upstream failure
        raise HTTPException(status_code=502, detail=f"MOEX error: {exc}") from exc


@app.get("/api/issuers/{issuer_id}/bonds")
def list_bonds(
    issuer_id: int, session: Session = Depends(get_session)
) -> list[dict[str, Any]]:
    return [credit_service.bond_dict(b) for b in repositories.list_bonds(session, issuer_id)]


@app.post("/api/issuers/{issuer_id}/bonds")
def add_bond(
    issuer_id: int,
    request: AddBondRequest,
    session: Session = Depends(get_session),
    client: MoexClient = Depends(get_moex),
) -> dict[str, Any]:
    if repositories.get_issuer(session, issuer_id) is None:
        raise HTTPException(status_code=404, detail=f"Unknown issuer {issuer_id}")
    try:
        bond = credit_service.add_bond(session, client, issuer_id, request.isin)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except Exception as exc:  # noqa: BLE001 - MOEX failure
        raise HTTPException(status_code=502, detail=f"MOEX error: {exc}") from exc
    return credit_service.bond_dict(bond)


@app.delete("/api/bonds/{isin}")
def delete_bond(isin: str, session: Session = Depends(get_session)) -> dict[str, str]:
    repositories.delete_bond(session, isin)
    return {"status": "deleted"}


@app.get("/api/bonds/{isin}/marketdata")
def bond_marketdata(
    isin: str,
    on: date | None = Query(default=None, alias="date"),
    session: Session = Depends(get_session),
    client: MoexClient = Depends(get_moex),
) -> dict[str, Any]:
    try:
        record = credit_service.fetch_market(session, client, isin, on or date.today())
    except Exception as exc:  # noqa: BLE001 - MOEX failure
        raise HTTPException(status_code=502, detail=f"MOEX error: {exc}") from exc
    return credit_service.market_dict(record)


# --- editable market prices ------------------------------------------------


@app.get("/api/issuers/{issuer_id}/prices")
def issuer_prices(
    issuer_id: int,
    on: date | None = Query(default=None, alias="date"),
    refresh: bool = Query(default=False),
    session: Session = Depends(get_session),
    client: MoexClient = Depends(get_moex),
) -> dict[str, Any]:
    try:
        return credit_service.issuer_prices(
            session, client, issuer_id, on or date.today(), force=refresh
        )
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.patch("/api/bonds/{isin}/prices")
def patch_price(
    isin: str,
    request: PriceOverride,
    on: date | None = Query(default=None, alias="date"),
    session: Session = Depends(get_session),
) -> dict[str, Any]:
    try:
        return credit_service.set_override(
            session, isin, on or date.today(), request.override_clean, request.included
        )
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


# --- credit curve ---------------------------------------------------------


@app.post("/api/credit-curve")
def credit_curve_endpoint(
    request: CreditCurveRequest,
    session: Session = Depends(get_session),
    client: MoexClient = Depends(get_moex),
) -> dict[str, Any]:
    try:
        return credit_service.credit_curve(
            session, client, request.issuer_id, request.rate_curve_id, request.trade_date
        )
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.get("/api/credit-curve.xlsx")
def credit_curve_excel(
    issuer_id: int = Query(),
    rate_curve_id: int = Query(),
    on: date | None = Query(default=None, alias="trade_date"),
    session: Session = Depends(get_session),
    client: MoexClient = Depends(get_moex),
) -> Response:
    try:
        ctx = credit_service.compute_credit(session, client, issuer_id, rate_curve_id, on)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    content = excel.build_workbook(ctx)
    stem = f"credit_{ctx.issuer_name}_{ctx.trade.isoformat()}".replace(" ", "_")
    return Response(
        content=content,
        media_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
        headers={"Content-Disposition": f'attachment; filename="{stem}.xlsx"'},
    )


# --- CDS pricing -----------------------------------------------------------


@app.get("/api/cds-pricing")
def cds_pricing(
    rate_curve_id: int = Query(),
    issuer_id: int | None = Query(default=None),
    on: date | None = Query(default=None, alias="trade_date"),
    refresh: bool = Query(default=False),
    session: Session = Depends(get_session),
    client: MoexClient = Depends(get_moex),
) -> dict[str, Any]:
    try:
        return cds_service.price_all_cds(
            session, client, rate_curve_id, on or date.today(), issuer_id, refresh
        )
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.get("/api/cds-pricing/{issuer_id}/{tenor_months}/breakdown")
def cds_pricing_breakdown(
    issuer_id: int,
    tenor_months: int,
    rate_curve_id: int = Query(),
    on: date | None = Query(default=None, alias="trade_date"),
    session: Session = Depends(get_session),
    client: MoexClient = Depends(get_moex),
) -> dict[str, Any]:
    try:
        return cds_service.cds_detail(
            session, client, issuer_id, rate_curve_id, on or date.today(), tenor_months
        )
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.get("/api/cds-pricing.xlsx")
def cds_pricing_excel(
    rate_curve_id: int = Query(),
    issuer_id: int | None = Query(default=None),
    on: date | None = Query(default=None, alias="trade_date"),
    session: Session = Depends(get_session),
    client: MoexClient = Depends(get_moex),
) -> Response:
    try:
        data = cds_service.price_all_cds(
            session, client, rate_curve_id, on or date.today(), issuer_id
        )
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "CDS Prices"
    headers = [
        "Issuer", "Tenor (m)", "Expiry", "Par spread (bp)", "Upfront",
        "Rebate", "Net premium", "DV01", "Credit DV01",
    ]
    for col, text in enumerate(headers, 1):
        cell = ws.cell(1, col, text)
        cell.font = Font(bold=True)
    for i, row in enumerate(data["results"], 2):
        ws.cell(i, 1, row.get("issuer_name", ""))
        ws.cell(i, 2, row["tenor_months"])
        ws.cell(i, 3, row.get("expiry_date", ""))
        ws.cell(i, 4, row["par_spread"])
        ws.cell(i, 5, row["upfront"])
        ws.cell(i, 6, row["rebate"])
        ws.cell(i, 7, row["net_premium"])
        ws.cell(i, 8, row["dv01"])
        ws.cell(i, 9, row["credit_dv01"])
    buf = io.BytesIO()
    wb.save(buf)
    stem = f"cds_{on or date.today()}"
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{stem}.xlsx"'},
    )


@app.get("/api/cds-pricing/{issuer_id}/{tenor_months}/breakdown.xlsx")
def cds_breakdown_excel(
    issuer_id: int,
    tenor_months: int,
    rate_curve_id: int = Query(),
    on: date | None = Query(default=None, alias="trade_date"),
    session: Session = Depends(get_session),
    client: MoexClient = Depends(get_moex),
) -> Response:
    try:
        detail = cds_service.cds_detail(
            session, client, issuer_id, rate_curve_id, on or date.today(), tenor_months
        )
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    _x = Font(bold=True)
    rows = [
        ("Issuer", detail["issuer_name"]),
        ("Tenor (m)", detail["tenor_months"]),
        ("Expiry", detail["expiry_date"]),
        ("Par spread (bp)", detail["par_spread"]),
        ("Upfront", detail["upfront"]),
        ("Rebate", detail["rebate"]),
        ("Net premium", detail["net_premium"]),
        ("DV01", detail["dv01"]),
        ("Credit DV01", detail["credit_dv01"]),
        ("Protection leg", detail["protection_leg"]),
        ("Premium leg", detail["premium_leg"]),
        ("RPV01", detail["rpv01"]),
        ("Recovery rate", detail.get("recovery_rate", "")),
        ("Base date", detail.get("base_date", "")),
    ]
    for i, (k, v) in enumerate(rows, 1):
        ws.cell(i, 1, k).font = _x
        ws.cell(i, 2, v)

    cf = wb.create_sheet("Cashflows")
    ch = ["Date", "Year fraction", "DF", "Survival", "Premium PV", "Protection PV"]
    for col, text in enumerate(ch, 1):
        cf.cell(1, col, text).font = _x
    for i, row in enumerate(detail.get("breakdown", []), 2):
        cf.cell(i, 1, row.get("date", ""))
        cf.cell(i, 2, row.get("year_fraction", ""))
        cf.cell(i, 3, row.get("df", ""))
        cf.cell(i, 4, row.get("survival", ""))
        cf.cell(i, 5, row.get("premium_pv", ""))
        cf.cell(i, 6, row.get("protection_pv", ""))

    if detail.get("discount_nodes"):
        ir = wb.create_sheet("IR Curve")
        for col, text in enumerate(["Date", "Years", "Zero rate %"], 1):
            ir.cell(1, col, text).font = _x
        for i, n in enumerate(detail["discount_nodes"], 2):
            ir.cell(i, 1, n.get("date", ""))
            ir.cell(i, 2, n.get("years", ""))
            ir.cell(i, 3, n.get("zero_rate_pct", ""))

    if detail.get("hazard_nodes"):
        hz = wb.create_sheet("Hazard Curve")
        for col, text in enumerate(["Date", "Years", "Hazard %"], 1):
            hz.cell(1, col, text).font = _x
        for i, n in enumerate(detail["hazard_nodes"], 2):
            hz.cell(i, 1, n.get("date", ""))
            hz.cell(i, 2, n.get("years", ""))
            hz.cell(i, 3, n.get("hazard_pct", ""))

    buf = io.BytesIO()
    wb.save(buf)
    name = f"cds_detail_{detail['issuer_name']}_{tenor_months}m".replace(" ", "_")
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{name}.xlsx"'},
    )


# --- data management --------------------------------------------------------


@app.get("/api/data/rate-indices")
def list_rate_indices(session: Session = Depends(get_session)) -> list[dict[str, Any]]:
    return [data_service.rate_index_dict(r) for r in repo.list_rate_indices(session)]


@app.post("/api/data/rate-indices")
def create_rate_index(
    request: dict[str, Any],
    session: Session = Depends(get_session),
) -> dict[str, Any]:
    ri = repo.upsert_rate_index(
        session,
        None,
        str(request.get("name", "")),
        str(request.get("currency", "")),
        str(request.get("day_count", "")),
        int(request.get("spot_lag", 0)),
        int(request.get("payment_lag", 0)),
        str(request.get("fixed_frequency", "")),
        str(request.get("business_day_convention", "")),
    )
    return data_service.rate_index_dict(ri)


@app.patch("/api/data/rate-indices/{index_id}")
def patch_rate_index(
    index_id: int,
    request: dict[str, Any],
    session: Session = Depends(get_session),
) -> dict[str, Any]:
    existing = repo.get_rate_index(session, index_id)
    if existing is None:
        raise HTTPException(status_code=404, detail="Unknown rate index")
    ri = repo.upsert_rate_index(
        session,
        index_id,
        str(request.get("name", existing.name)),
        str(request.get("currency", existing.currency)),
        str(request.get("day_count", existing.day_count)),
        int(request.get("spot_lag", existing.spot_lag)),
        int(request.get("payment_lag", existing.payment_lag)),
        str(request.get("fixed_frequency", existing.fixed_frequency)),
        str(request.get("business_day_convention", existing.business_day_convention)),
    )
    return data_service.rate_index_dict(ri)


@app.delete("/api/data/rate-indices/{index_id}")
def delete_rate_index(
    index_id: int, session: Session = Depends(get_session)
) -> dict[str, str]:
    repo.delete_rate_index(session, index_id)
    return {"status": "deleted"}


@app.get("/api/data/calendars")
def list_calendars(session: Session = Depends(get_session)) -> list[dict[str, Any]]:
    return [data_service.calendar_dict(c) for c in repo.list_calendars(session)]


@app.post("/api/data/calendars")
def create_calendar(
    request: dict[str, Any],
    session: Session = Depends(get_session),
) -> dict[str, Any]:
    cal = repo.create_calendar(
        session,
        str(request.get("code", "")),
        str(request.get("description", "")),
        str(request.get("currency", "")),
        bool(request.get("is_default_for_cds", False)),
        bool(request.get("is_default_for_ois", False)),
    )
    return data_service.calendar_dict(cal)


@app.delete("/api/data/calendars/{calendar_id}")
def delete_calendar(
    calendar_id: int, session: Session = Depends(get_session)
) -> dict[str, str]:
    repo.delete_calendar(session, calendar_id)
    return {"status": "deleted"}


@app.get("/api/data/calendars/{calendar_id}/holidays")
def list_holidays(
    calendar_id: int, session: Session = Depends(get_session)
) -> list[dict[str, Any]]:
    return [data_service.holiday_dict(h) for h in repo.list_holidays(session, calendar_id)]


@app.post("/api/data/calendars/{calendar_id}/holidays")
def add_holiday(
    calendar_id: int,
    request: dict[str, Any],
    session: Session = Depends(get_session),
) -> dict[str, Any]:
    h = repo.add_holiday(session, calendar_id, date.fromisoformat(str(request["date"])))
    return data_service.holiday_dict(h)


@app.delete("/api/data/calendars/{calendar_id}/holidays/{holiday_id}")
def delete_holiday(
    calendar_id: int, holiday_id: int, session: Session = Depends(get_session)
) -> dict[str, str]:
    repo.delete_holiday(session, holiday_id)
    return {"status": "deleted"}


@app.patch("/api/data/calendars/{calendar_id}")
def patch_calendar(
    calendar_id: int,
    request: dict[str, Any],
    session: Session = Depends(get_session),
) -> dict[str, Any]:
    cal = repo.get_calendar(session, calendar_id)
    if cal is None:
        raise HTTPException(status_code=404, detail="Unknown calendar")
    if "code" in request:
        cal.code = request["code"]
    if "description" in request:
        cal.description = request["description"]
    if "currency" in request:
        cal.currency = request["currency"].upper()
    if "is_default_for_cds" in request:
        cal.is_default_for_cds = bool(request["is_default_for_cds"])
    if "is_default_for_ois" in request:
        cal.is_default_for_ois = bool(request["is_default_for_ois"])
    session.commit()
    session.refresh(cal)
    return data_service.calendar_dict(cal)


@app.get("/api/data/rate-indices/export")
def export_rate_indices(session: Session = Depends(get_session)) -> Response:
    mapping = repo.build_rate_index_mapping(session)
    return Response(
        content=json.dumps(mapping, indent=2),
        media_type="application/json",
        headers={"Content-Disposition": "attachment; filename=rate_indices.json"},
    )


@app.get("/api/data/calendars/{calendar_id}/export")
def export_calendar(
    calendar_id: int, session: Session = Depends(get_session)
) -> Response:
    cal = repo.get_calendar(session, calendar_id)
    if cal is None:
        raise HTTPException(status_code=404, detail="Unknown calendar")
    dates = sorted(h.date.isoformat() for h in cal.dates)
    payload = {cal.currency: dates}
    name = f"calendar_{cal.code}.json"
    return Response(
        content=json.dumps(payload, indent=2),
        media_type="application/json",
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )


@app.get("/")
def index() -> FileResponse:
    return FileResponse(_FRONTEND / "index.html")


app.mount("/", StaticFiles(directory=_FRONTEND), name="static")
